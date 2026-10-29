"""Shared helpers for tests."""

from __future__ import annotations

import io


def make_draft(**overrides):
    base = {
        "cv_draft_version": "0.1",
        # These fixtures exercise rendering, not provenance. Declaring the mode
        # explicitly keeps the export evidence contract closed by default.
        "evidence_mode": "unverified",
        "target_job": {"title": "AI Engineer", "company": "Acme"},
        "candidate": {
            "full_name": "Jane Example",
            "email": "jane@example.com",
            "phone": "000 000 000",
            "location": "Madrid, Spain",
            "links": ["https://example.com/portfolio", "https://example.com/profile"],
        },
        "sections": [
            {
                "section_id": "section_summary",
                "kind": "summary",
                "title": "Professional Summary",
                "items": [
                    {
                        "item_id": "item_1",
                        "text": "Experienced engineer with expertise in Python.",
                        "approved_fact_ids": ["fact_1"],
                        "review_status": "approved",
                    }
                ],
            },
            {
                "section_id": "section_experience",
                "kind": "experience",
                "title": "Experience",
                "items": [
                    {
                        "item_id": "item_2",
                        "text": "Data Engineer at Acme",
                        "approved_fact_ids": ["fact_2"],
                        "review_status": "approved",
                        "heading": "Data Engineer",
                        "subheading": "Acme",
                        "period": "2020 - 2023",
                        "location": "Madrid",
                    }
                ],
            },
            {
                "section_id": "section_skills",
                "kind": "skills",
                "title": "Skills",
                "items": [
                    {"item_id": "item_3", "text": "Python", "approved_fact_ids": ["fact_3"], "review_status": "approved"},
                    {"item_id": "item_4", "text": "SQL", "approved_fact_ids": ["fact_4"], "review_status": "approved"},
                ],
            },
        ],
        "warnings": [],
        "unmet_requirements": ["Kubernetes"],
    }
    base.update(overrides)
    return base


def pdf_text(path_or_bytes):
    from pypdf import PdfReader

    if isinstance(path_or_bytes, (bytes, bytearray)):
        r = PdfReader(io.BytesIO(path_or_bytes))
    else:
        r = PdfReader(path_or_bytes)
    return "\n".join((p.extract_text() or "") for p in r.pages)


def docx_text(path_or_bytes):
    import docx

    if isinstance(path_or_bytes, (bytes, bytearray)):
        d = docx.Document(io.BytesIO(path_or_bytes))
    else:
        d = docx.Document(path_or_bytes)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)