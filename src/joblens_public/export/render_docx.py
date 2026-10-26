"""python-docx-based DOCX renderer."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import CanonicalDraft
from .utils import link_label, normalize_url

ACCENT = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0x44, 0x44, 0x44)


def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.08
    # East asian font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Calibri")

    for section in doc.sections:
        section.top_margin = Cm(1.35)
        section.bottom_margin = Cm(1.35)
        section.left_margin = Cm(1.35)
        section.right_margin = Cm(1.35)


def _add_hyperlink(paragraph, url: str, text: str | None = None) -> None:
    url = normalize_url(url)
    display = text if text else link_label(url)
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0F766E")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = display
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _para(doc, text: str = "", *, bold=False, italic=False, size=10, color=None, align=None, space_after=3, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return p


def _render_contact(doc, draft: CanonicalDraft) -> None:
    c = draft.candidate
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    first = True
    def add_sep():
        nonlocal first
        if not first:
            r = p.add_run("  ·  ")
            r.font.size = Pt(9)
            r.font.color.rgb = GREY
        first = False

    if c.location:
        add_sep()
        r = p.add_run(c.location)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    if c.email:
        add_sep()
        r = p.add_run(c.email)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        # make clickable
        # Rebuild as hyperlink instead: easier to append hyperlink after run? We'll just add hyperlink.
    # Simpler: build hyperlinks separately for email/links.
    # Remove runs above? We'll redo cleanly below.
    # Actually rebuild paragraph from scratch:
    # Clear
    for child in list(p._p):
        p._p.remove(child)
    # Rebuild with hyperlinks
    first = True
    def sep():
        nonlocal first
        if not first:
            r = p.add_run("  ·  ")
            r.font.size = Pt(9)
            r.font.color.rgb = GREY
        first = False
    if c.location:
        sep()
        r = p.add_run(c.location)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    if c.email:
        sep()
        _add_hyperlink(p, "mailto:" + c.email, c.email)
        # adjust size/color not trivial; set run props via last run
        last = p.runs[-1] if p.runs else None
        if last is not None:
            last.font.size = Pt(9)
    if c.phone:
        sep()
        r = p.add_run(c.phone)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    for link in c.links:
        sep()
        _add_hyperlink(p, link, link_label(link))
        last = p.runs[-1] if p.runs else None
        if last is not None:
            last.font.size = Pt(9)


def render_docx(draft: CanonicalDraft, destination: Any) -> None:
    doc = Document()
    _set_base_style(doc)

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_r = name_p.add_run(draft.candidate.full_name)
    name_r.bold = True
    name_r.font.size = Pt(17)
    name_p.paragraph_format.space_after = Pt(2)

    _render_contact(doc, draft)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    for section in draft.sections:
        # Section title
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(9)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(section.title)
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = ACCENT
        # bottom border
        pPr = h._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "2")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "0F766E")
        pbdr.append(bottom)
        pPr.append(pbdr)

        kind = (section.kind or "").lower()
        if kind in ("skills", "languages", "certifications"):
            names = [ (item.text or "").strip() for item in section.items if (item.text or "").strip() ]
            if names:
                p = doc.add_paragraph()
                p.add_run(", ".join(names))
            continue
        if kind == "summary":
            for item in section.items:
                if (item.text or "").strip():
                    doc.add_paragraph(item.text.strip())
            continue
        if kind == "education":
            for item in section.items:
                if (item.heading or "").strip() and (item.heading or "").strip().lower() != (item.text or "").strip().lower():
                    p = doc.add_paragraph()
                    r = p.add_run(item.heading.strip())
                    r.bold = True
                    p.paragraph_format.space_after = Pt(0)
                if (item.text or "").strip():
                    doc.add_paragraph(item.text.strip())
                if item.period:
                    p = doc.add_paragraph(item.period.strip())
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.font.color.rgb = GREY
                if item.location:
                    p = doc.add_paragraph(item.location.strip())
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.font.color.rgb = GREY
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # experience/projects/generic
        for item in section.items:
            # heading line
            if item.heading or item.subheading or item.period:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = True
                if item.heading:
                    r = p.add_run(item.heading)
                    r.bold = True
                if item.subheading and item.subheading != item.heading:
                    if item.heading:
                        p.add_run("  ")
                    r = p.add_run(item.subheading)
                    r.font.color.rgb = GREY
                if item.period:
                    # add tab to right
                    p.add_run("\t" + (item.period or ""))
                # right align period via tab stop
                from docx.enum.text import WD_TAB_ALIGNMENT
                p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
            combined = " ".join(x for x in (item.heading, item.subheading) if x).strip().lower()
            text = (item.text or "").strip()
            if text and text.lower() != combined and text.lower() != (item.heading or "").strip().lower():
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(text)
            if item.location:
                p = doc.add_paragraph(item.location.strip())
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = GREY

    if hasattr(destination, "write"):
        doc.save(destination)
    else:
        doc.save(str(destination))


def render_docx_bytes(draft: CanonicalDraft) -> bytes:
    buf = io.BytesIO()
    render_docx(draft, buf)
    return buf.getvalue()