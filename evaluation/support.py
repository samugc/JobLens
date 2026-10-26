"""Small helpers for the synthetic evaluation harness."""

from __future__ import annotations

import http.client
import io
import json
import re
import threading
from dataclasses import dataclass
from typing import Any

from joblens_public.data import load_demo_data
from joblens_public.webapp import create_server
from tests.cv_builders import make_docx, make_pdf
from tests.cv_http_helpers import build_multipart


@dataclass(frozen=True)
class Document:
    filename: str
    kind: str  # "pdf" or "docx"
    text: str
    pages: int = 1

    def build(self) -> bytes:
        lines = [line for line in self.text.splitlines()]
        if self.kind == "pdf":
            if self.pages <= 1:
                return make_pdf([lines])
            # Split lines across the requested number of pages.
            # Split at contiguous boundaries so a section is not scattered
            # across pages. The evaluation targets content preservation, not
            # adversarial page interleaving.
            size = max(1, (len(lines) + self.pages - 1) // self.pages)
            chunks = [lines[index : index + size] for index in range(0, len(lines), size)]
            return make_pdf(chunks)
        if self.kind == "docx":
            return make_docx([(line, None) for line in lines])
        raise ValueError(f"unsupported document kind {self.kind!r}")


class LocalServer:
    def __init__(self) -> None:
        candidates, jobs = load_demo_data()
        self.server = create_server(candidates, jobs, 0)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)

    def __enter__(self) -> "LocalServer":
        self.thread.start()
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=3)

    @property
    def port(self) -> int:
        return self.server.server_port

    def request(
        self,
        method: str,
        path: str,
        body: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, dict[str, str], bytes]:
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=30)
        try:
            connection.request(method, path, body=body, headers=headers or {})
            response = connection.getresponse()
            raw = response.read()
            return (
                response.status,
                {name.casefold(): value for name, value in response.getheaders()},
                raw,
            )
        finally:
            connection.close()

    def post_json(self, path: str, payload: Any) -> tuple[int, dict[str, str], bytes]:
        headers = {
            "Content-Type": "application/json",
            "Origin": f"http://127.0.0.1:{self.port}",
            "X-JobLens-Request": "1",
            "Accept": "application/json",
        }
        return self.request("POST", path, json.dumps(payload).encode("utf-8"), headers)

    def post_files(self, documents: list[Document]) -> tuple[int, dict[str, str], bytes]:
        files = [("files", document.filename, document.build()) for document in documents]
        body, boundary = build_multipart(files)
        headers = {
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "Origin": f"http://127.0.0.1:{self.port}",
            "X-JobLens-Request": "1",
            "Accept": "application/json",
        }
        return self.request("POST", "/api/cv/extractions", body, headers)


def json_payload(raw: bytes) -> dict[str, Any]:
    value = json.loads(raw.decode("utf-8"))
    if not isinstance(value, dict):
        raise ValueError("expected a JSON object")
    return value


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def docx_text(data: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def normalise(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().casefold()


def contains(text: str, needle: str) -> bool:
    return normalise(needle) in normalise(text)


def significant_tokens(text: str) -> set[str]:
    return {
        token
        for token in re.findall(r"[a-z0-9][a-z0-9+#./-]{2,}", normalise(text))
        if token not in {"the", "and", "with", "for", "from", "that", "this"}
    }


def token_jaccard(left: str, right: str) -> float:
    a = significant_tokens(left)
    b = significant_tokens(right)
    if not a and not b:
        return 1.0
    return len(a & b) / len(a | b)
